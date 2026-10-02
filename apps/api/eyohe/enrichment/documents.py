"""Public document intelligence: metadata and text extraction for PDF, DOCX, XLSX, TXT, HTML."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any

from dateutil import parser as dateparser


@dataclass
class DocumentInfo:
    mime_type: str
    title: str = ""
    author: str = ""
    creator: str = ""
    producer: str = ""
    subject: str = ""
    created_at: datetime | None = None
    modified_at: datetime | None = None
    pages: int | None = None
    text: str = ""
    metadata: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def _pdf_date(v: object) -> datetime | None:
    if not v:
        return None
    s = str(v)
    m = re.match(r"D:(\d{4})(\d{2})?(\d{2})?(\d{2})?(\d{2})?(\d{2})?", s)
    try:
        if m:
            y, mo, dd, hh, mm, ss = (int(x) if x else (1 if i < 3 else 0) for i, x in enumerate(m.groups()))
            return datetime(y, mo, dd, hh, mm, ss, tzinfo=UTC)
        parsed = dateparser.parse(s)
        return parsed if parsed.tzinfo else parsed.replace(tzinfo=UTC)
    except (ValueError, TypeError, OverflowError):
        return None


def extract_document(content: bytes, mime_type: str, *, max_chars: int = 200_000) -> DocumentInfo:
    mt = mime_type.split(";")[0].strip().lower()
    info = DocumentInfo(mime_type=mt)
    try:
        if mt == "application/pdf" or content[:5] == b"%PDF-":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                info.warnings.append("PDF is encrypted; metadata/text not extracted (no bypass attempted).")
                return info
            meta: dict[str, Any] = dict(reader.metadata or {})
            info.title = str(meta.get("/Title", "") or "")
            info.author = str(meta.get("/Author", "") or "")
            info.creator = str(meta.get("/Creator", "") or "")
            info.producer = str(meta.get("/Producer", "") or "")
            info.subject = str(meta.get("/Subject", "") or "")
            info.created_at = _pdf_date(meta.get("/CreationDate"))
            info.modified_at = _pdf_date(meta.get("/ModDate"))
            info.pages = len(reader.pages)
            chunks = []
            total = 0
            for page in reader.pages[:200]:
                t = page.extract_text() or ""
                chunks.append(t)
                total += len(t)
                if total > max_chars:
                    break
            info.text = "\n".join(chunks)[:max_chars]
            info.metadata = {k.lstrip("/"): str(v)[:300] for k, v in dict(meta).items()}
        elif mt in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",) or (
            content[:2] == b"PK" and b"word/" in content[:4000]
        ):
            import docx

            d = docx.Document(io.BytesIO(content))
            cp = d.core_properties
            info.title, info.author, info.subject = cp.title or "", cp.author or "", cp.subject or ""
            info.created_at = cp.created.replace(tzinfo=UTC) if cp.created and cp.created.tzinfo is None else cp.created
            info.modified_at = (
                cp.modified.replace(tzinfo=UTC) if cp.modified and cp.modified.tzinfo is None else cp.modified
            )
            info.metadata = {
                "last_modified_by": cp.last_modified_by or "",
                "revision": cp.revision,
                "category": cp.category or "",
                "comments": (cp.comments or "")[:300],
            }
            info.text = "\n".join(p.text for p in d.paragraphs)[:max_chars]
        elif mt in ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",):
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            p = wb.properties
            info.title, info.author = p.title or "", p.creator or ""
            info.created_at = p.created.replace(tzinfo=UTC) if p.created and p.created.tzinfo is None else p.created
            info.modified_at = (
                p.modified.replace(tzinfo=UTC) if p.modified and p.modified.tzinfo is None else p.modified
            )
            info.metadata = {"last_modified_by": p.lastModifiedBy or "", "sheets": wb.sheetnames[:50]}
            rows = []
            total = 0
            for ws in wb.worksheets[:10]:
                for row in ws.iter_rows(values_only=True):
                    line = " | ".join("" if c is None else str(c) for c in row)
                    rows.append(line)
                    total += len(line)
                    if total > max_chars:
                        break
            info.text = "\n".join(rows)[:max_chars]
        elif mt.startswith("text/html"):
            from eyohe.collectors.web_collector import html_to_text

            title, text, meta = html_to_text(content.decode("utf-8", "replace"))
            info.title, info.text, info.metadata = title, text[:max_chars], meta
            info.author = str(meta.get("author", ""))
        elif mt.startswith("text/"):
            info.text = content.decode("utf-8", "replace")[:max_chars]
        else:
            info.warnings.append(f"unsupported document type {mt}")
    except Exception as exc:
        info.warnings.append(f"extraction failed: {type(exc).__name__}: {exc}"[:300])
    return info
