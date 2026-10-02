"""Render ReportData to Markdown, HTML, PDF (WeasyPrint → headless Chromium fallback) and DOCX."""

from __future__ import annotations

import asyncio
import shutil
import subprocess
import tempfile
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape

from eyohe.core.config import get_settings
from eyohe.core.errors import ConfigurationError
from eyohe.core.logging import get_logger
from eyohe.reporting.builder import ReportData

log = get_logger("reporting")
_env = Environment(
    loader=FileSystemLoader(Path(__file__).parent / "templates"), autoescape=select_autoescape(["html", "j2"])
)


def render_html(d: ReportData) -> str:
    return _env.get_template("report.html.j2").render(d=d)


def render_markdown(d: ReportData) -> str:
    lines: list[str] = []
    a = lines.append
    a("# EYOHE OSINT — INVESTIGATION REPORT")
    a("")
    if d.case.is_demo:
        a("> **DEMO DATA** — every entity in this report is fictional.")
        a("")
    a(f"**Investigation:** {d.case.name}  ")
    a(f"**Case ID:** `{d.case.display_id}`  ")
    a(f"**Date:** {d.generated_at.strftime('%Y-%m-%d %H:%M UTC')}  ")
    a(f"**Analyst:** {d.analyst}  ")
    a(f"**Classification:** {d.classification}")
    a("")
    a("## Executive Summary")
    a(
        f"{d.stats['sources']} sources, {d.stats['evidence']} evidence records, {d.stats['entities']} entities, {d.stats['relationships']} relationships, {d.stats['findings']} findings."
    )
    for f in d.findings[:6]:
        a(f"- **[{f['confidence']}]** {f['title']} (`{f['id']}`)")
    a("")
    a("## Investigation Objective")
    a(d.case.objective or "_No objective recorded._")
    a("")
    a("## Targets")
    a("| Type | Value | Primary |")
    a("|---|---|---|")
    for t in d.targets:
        a(f"| {t.type} | `{t.normalized_value}` | {'yes' if t.is_primary else ''} |")
    a("")
    a("## Key Findings")
    a("_Confidence is computed from attached evidence, not asserted by an AI._")
    for f in d.findings:
        a("")
        a(f"### {f['id']} — {f['title']}")
        a(f"**Assessment:** {f['confidence']} · category {f['category']} · review {f['review_state']}  ")
        a(f"**Claim:** {f['claim']}  ")
        if f["assessment"]:
            a(f"**Analyst assessment:** {f['assessment']}  ")
        if f["rationale"].get("reasons"):
            a(f"**Why this confidence:** {' '.join(f['rationale']['reasons'])}  ")
        a(
            "**Supporting evidence:** "
            + (
                ", ".join(
                    f"`{e['id']}`" + (f" ({e['source_id']})" if e.get("source_id") else "") for e in f["supporting"]
                )
                or "none"
            )
        )
        if f["contradicting"]:
            a(
                "**Contradicting evidence (CONFLICTING SOURCES):** "
                + ", ".join(f"`{e['id']}`" for e in f["contradicting"])
            )
    if not d.findings:
        a("_None._")
    a("")
    a("## Entity Analysis")
    a("| ID | Type | Value | Confidence | Evidence |")
    a("|---|---|---|---|---|")
    for e in d.entities[:200]:
        a(
            f"| `{e.display_id}` | {e.type} | {e.value}{' **(target)**' if e.is_target else ''} | {e.confidence} | {e.source_count} |"
        )
    a("")
    a("## Relationship Graph")
    a("| ID | Source | Relationship | Target | Confidence | Evidence |")
    a("|---|---|---|---|---|---|")
    for r in d.relationships[:300]:
        a(
            f"| `{r['id']}` | {r['source']} | {r['type']} | {r['target']} | {r['confidence']} | {', '.join(r['evidence_ids'])} |"
        )
    a("")
    a("## Timeline")
    for tl in d.timeline:
        a(f"- `{tl.occurred_at.strftime('%Y-%m-%d')}` — {tl.title} [{tl.confidence}]")
    if not d.timeline:
        a("_No dated events._")
    a("")
    a("## Technical Findings")
    for ev in [x for x in d.evidence if x["type"] in ("TECHNICAL_RECORD", "ARCHIVE_SNAPSHOT", "DOCUMENT")][:150]:
        a(
            f"- `{ev['id']}` {ev['claim']}"
            + (f" — source `{ev['source']['id']}` {ev['source']['url']}" if ev["source"] else "")
        )
    a("")
    a("## Social / Public-Source Findings")
    a("_Statements by third parties are reported as statements, never as facts; allegations are labelled._")
    for ev in [x for x in d.evidence if x["type"] in ("PUBLIC_STATEMENT", "ALLEGATION", "DIRECT_STATEMENT")][:150]:
        a(
            f"- `{ev['id']}` **{ev['language']}** — {ev['claim']}"
            + (f' — "{ev["excerpt"][:200]}"' if ev["excerpt"] else "")
            + (f" — {ev['source']['url']}" if ev["source"] else "")
        )
    a("")
    a("## Source Analysis")
    a("| ID | Title | Type | Tier | Published | Collected | URL |")
    a("|---|---|---|---|---|---|---|")
    for s in d.sources:
        a(
            f"| `{s.display_id}` | {s.title[:60].replace('|', '/')} | {s.source_type} | {s.tier} | {s.published_at.strftime('%Y-%m-%d') if s.published_at else '—'} | {s.collected_at.strftime('%Y-%m-%d')} | {s.url} |"
        )
    a("")
    a("## Evidence Index")
    a("| ID | Type | Claim | Source | Confidence | Review |")
    a("|---|---|---|---|---|---|")
    for ev in d.evidence:
        a(
            f"| `{ev['id']}` | {ev['type']} | {ev['claim'].replace('|', '/')} | {ev['source']['id'] if ev['source'] else ev['collector']} | {ev['confidence']} | {ev['review_state']} |"
        )
    a("")
    a("## Limitations")
    for l_ in d.limitations:
        a(f"- {l_}")
    a("")
    a("## Methodology")
    a(
        "Plan → collectors → sources (URL, timestamp, hash) → evidence (cited, quote-validated) → entities/relationships (evidence-backed) → rule-based confidence → analyst review. Public sources only."
    )
    for inv in d.investigations:
        a("")
        a(f"**{inv.display_id} task log**")
        for task in d.tasks[inv.id]:
            a(
                f"- {task.title}: {task.status}"
                + (f" — {task.error[:100]}" if task.error else "")
                + f" — _{task.rationale}_"
            )
    a("")
    a("## References")
    for i, s in enumerate(d.sources, 1):
        a(
            f"{i}. `{s.display_id}` {s.title[:100]}{' — ' + s.publisher if s.publisher else ''}. {s.url} (collected {s.collected_at.strftime('%Y-%m-%d')})"
        )
    a("")
    a(f"_Generated by Eyohe OSINT v{d.version} · {d.case.display_id} · {d.classification}_")
    return "\n".join(lines)


async def render_pdf(html: str, out_path: Path) -> str:
    """Return the engine used. WeasyPrint if importable, else Chromium headless."""
    try:
        from weasyprint import HTML

        await asyncio.to_thread(lambda: HTML(string=html, base_url=str(out_path.parent)).write_pdf(str(out_path)))
        return "weasyprint"
    except ImportError:
        pass
    except Exception as exc:
        log.warning("weasyprint_failed", error=str(exc)[:200])
    chromium = get_settings().chromium_path
    binary = (
        chromium
        if Path(chromium).exists()
        else (shutil.which("chromium") or shutil.which("chromium-browser") or shutil.which("google-chrome"))
    )
    if not binary:
        raise ConfigurationError(
            "PDF generation needs WeasyPrint (uv sync --extra pdf) or Chromium (CHROMIUM_PATH). Neither is available."
        )
    with tempfile.TemporaryDirectory(prefix="eyohe-pdf-") as tmp:
        src = Path(tmp) / "report.html"
        src.write_text(html, encoding="utf-8")
        argv = [
            binary,
            "--headless=new",
            "--disable-gpu",
            "--no-sandbox",
            "--no-first-run",
            "--no-default-browser-check",
            f"--user-data-dir={tmp}/profile",
            "--no-pdf-header-footer",
            f"--print-to-pdf={out_path}",
            src.as_uri(),
        ]
        proc = await asyncio.create_subprocess_exec(*argv, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
        try:
            _, err = await asyncio.wait_for(proc.communicate(), timeout=120)
        except TimeoutError:
            proc.kill()
            raise ConfigurationError("Chromium PDF rendering timed out after 120s.") from None
        if proc.returncode != 0 or not out_path.exists():
            raise ConfigurationError(f"Chromium PDF rendering failed: {err.decode(errors='replace')[-300:]}")
    return "chromium"


def render_docx(d: ReportData, out_path: Path) -> None:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    def heading(text: str, level: int = 1) -> None:
        doc.add_heading(text, level=level)

    def table(headers: list[str], rows: list[list[str]]) -> None:
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = h
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = str(v)[:1000]

    p = doc.add_paragraph()
    run = p.add_run("EYOHE OSINT")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    doc.add_paragraph("Public Intelligence. Connected Evidence.")
    t = doc.add_paragraph()
    t.add_run("INVESTIGATION REPORT").bold = True
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if d.case.is_demo:
        doc.add_paragraph("DEMO DATA — every entity in this report is fictional.")
    for label, val in (
        ("Investigation", d.case.name),
        ("Case ID", d.case.display_id),
        ("Date", d.generated_at.strftime("%Y-%m-%d %H:%M UTC")),
        ("Analyst", d.analyst),
        ("Classification", d.classification),
    ):
        para = doc.add_paragraph()
        para.add_run(f"{label}: ").bold = True
        para.add_run(val)
    doc.add_page_break()  # type: ignore[no-untyped-call]
    heading("1. Executive Summary")
    doc.add_paragraph(
        f"{d.stats['sources']} sources, {d.stats['evidence']} evidence records, {d.stats['entities']} entities, {d.stats['relationships']} relationships, {d.stats['findings']} findings."
    )
    for f in d.findings[:6]:
        doc.add_paragraph(f"[{f['confidence']}] {f['title']} ({f['id']})", style="List Bullet")
    heading("2. Investigation Objective")
    doc.add_paragraph(d.case.objective or "No objective recorded.")
    heading("3. Targets")
    table(
        ["Type", "Value", "Primary"], [[t.type, t.normalized_value, "yes" if t.is_primary else ""] for t in d.targets]
    )
    heading("4. Key Findings")
    doc.add_paragraph("Confidence labels are computed from the attached evidence, not asserted by an AI.")
    for f in d.findings:
        heading(f"{f['id']} — {f['title']}", 2)
        doc.add_paragraph(f"Assessment: {f['confidence']} · category {f['category']} · review {f['review_state']}")
        doc.add_paragraph(f"Claim: {f['claim']}")
        if f["assessment"]:
            doc.add_paragraph(f"Analyst assessment: {f['assessment']}")
        if f["rationale"].get("reasons"):
            doc.add_paragraph("Why this confidence: " + " ".join(f["rationale"]["reasons"]))
        doc.add_paragraph(
            "Supporting evidence: "
            + (
                ", ".join(e["id"] + (f" ({e['source_id']})" if e.get("source_id") else "") for e in f["supporting"])
                or "none"
            )
        )
        if f["contradicting"]:
            doc.add_paragraph(
                "Contradicting evidence (CONFLICTING SOURCES): " + ", ".join(e["id"] for e in f["contradicting"])
            )
    heading("5. Entity Analysis")
    table(
        ["ID", "Type", "Value", "Confidence", "Evidence"],
        [
            [e.display_id, e.type, e.value + (" (target)" if e.is_target else ""), e.confidence, str(e.source_count)]
            for e in d.entities[:200]
        ],
    )
    heading("6. Relationship Graph")
    table(
        ["ID", "Source", "Relationship", "Target", "Confidence", "Evidence"],
        [
            [r["id"], r["source"], r["type"], r["target"], r["confidence"], ", ".join(r["evidence_ids"])]
            for r in d.relationships[:300]
        ],
    )
    heading("7. Timeline")
    for t_ in d.timeline:
        doc.add_paragraph(f"{t_.occurred_at.strftime('%Y-%m-%d')} — {t_.title} [{t_.confidence}]", style="List Bullet")
    heading("8. Technical Findings")
    table(
        ["Evidence", "Record", "Source", "Collected"],
        [
            [e["id"], e["claim"], (e["source"]["url"] if e["source"] else ""), e["collected_at"].strftime("%Y-%m-%d")]
            for e in d.evidence
            if e["type"] in ("TECHNICAL_RECORD", "ARCHIVE_SNAPSHOT", "DOCUMENT")
        ][:150],
    )
    heading("9. Social / Public-Source Findings")
    doc.add_paragraph(
        "Statements by third parties are reported as statements, never as facts; allegations are labelled."
    )
    table(
        ["Evidence", "Language", "Statement", "Source"],
        [
            [
                e["id"],
                e["language"],
                e["claim"] + (f' — "{e["excerpt"][:200]}"' if e["excerpt"] else ""),
                (e["source"]["url"] if e["source"] else ""),
            ]
            for e in d.evidence
            if e["type"] in ("PUBLIC_STATEMENT", "ALLEGATION", "DIRECT_STATEMENT")
        ][:150],
    )
    heading("10. Source Analysis")
    table(
        ["ID", "Title", "Type", "Tier", "Published", "Collected", "URL"],
        [
            [
                s.display_id,
                s.title[:80],
                s.source_type,
                str(s.tier),
                s.published_at.strftime("%Y-%m-%d") if s.published_at else "—",
                s.collected_at.strftime("%Y-%m-%d"),
                s.url,
            ]
            for s in d.sources
        ],
    )
    heading("11. Evidence Index")
    table(
        ["ID", "Type", "Claim", "Source", "Confidence", "Review"],
        [
            [
                e["id"],
                e["type"],
                e["claim"],
                (e["source"]["id"] if e["source"] else e["collector"]),
                e["confidence"],
                e["review_state"],
            ]
            for e in d.evidence
        ],
    )
    heading("12. Limitations")
    for l_ in d.limitations:
        doc.add_paragraph(l_, style="List Bullet")
    heading("13. Methodology")
    doc.add_paragraph(
        "Plan → collectors → sources (URL, timestamp, hash) → evidence (cited, quote-validated) → entities/relationships (evidence-backed) → rule-based confidence → analyst review. Public sources only."
    )
    heading("14. References")
    for i, s in enumerate(d.sources, 1):
        doc.add_paragraph(
            f"{i}. {s.display_id} {s.title[:100]} — {s.url} (collected {s.collected_at.strftime('%Y-%m-%d')})"
        )
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = f"Eyohe OSINT v{d.version} · {d.case.display_id} · {d.classification} · generated {d.generated_at.strftime('%Y-%m-%d %H:%M UTC')}"
    doc.save(str(out_path))
